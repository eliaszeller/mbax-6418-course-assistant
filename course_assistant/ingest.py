from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from .models import SourceChunk


ACCEPTED = {".pdf", ".pptx", ".ppt", ".odp", ".docx", ".txt", ".md"}


def _id(path: Path, locator: str, kind: str) -> str:
    return hashlib.sha256(f"{path.resolve()}|{locator}|{kind}".encode()).hexdigest()[:20]


def _render_pdf(pdf: Path, image_dir: Path) -> list[Path]:
    image_dir.mkdir(parents=True, exist_ok=True)
    prefix = image_dir / "page"
    tool = shutil.which("pdftoppm")
    if not tool:
        return []
    subprocess.run([tool, "-png", "-r", "130", str(pdf), str(prefix)], check=True)
    return sorted(image_dir.glob("page-*.png"))


def _pptx_to_pdf(path: Path, work: Path) -> Path | None:
    tool = shutil.which("soffice") or shutil.which("libreoffice")
    if not tool:
        return None
    work.mkdir(parents=True, exist_ok=True)
    subprocess.run([tool, "--headless", "--convert-to", "pdf", "--outdir", str(work), str(path)], check=True)
    candidate = work / f"{path.stem}.pdf"
    return candidate if candidate.exists() else None


def _ingest_pdf_pages(path: Path, source_name: str, material: str, output_root: Path,
                      locator_name: str = "Page") -> list[SourceChunk]:
    asset_dir = output_root / "assets" / hashlib.sha1(str(path).encode()).hexdigest()[:12]
    reader = PdfReader(path)
    rendered = _render_pdf(path, asset_dir / "pages")
    chunks: list[SourceChunk] = []
    for number, page in enumerate(reader.pages, 1):
        locator = f"{locator_name} {number}"
        text = page.extract_text() or f"[{locator_name} with no extractable text]"
        image = str(rendered[number - 1]) if number <= len(rendered) else None
        chunks.append(SourceChunk(id=_id(path, locator, "text"), document=source_name, material=material,
                                  locator=locator, text=text, image_path=image))
        if image:
            chunks.append(SourceChunk(id=_id(path, locator, "visual"), document=source_name, material=material,
                                      locator=locator, text=f"Visual evidence from {locator}: {text}",
                                      image_path=image, kind="visual"))
    return chunks


def ingest_file(path: Path, output_root: Path) -> list[SourceChunk]:
    path = path.resolve()
    if path.suffix.lower() not in ACCEPTED:
        raise ValueError(f"Unsupported file: {path.name}. Accepted: {sorted(ACCEPTED)}")
    material = path.stem
    asset_dir = output_root / "assets" / hashlib.sha1(str(path).encode()).hexdigest()[:12]
    chunks: list[SourceChunk] = []

    if path.suffix.lower() == ".pptx":
        deck = Presentation(path)
        rendered: list[Path] = []
        pdf = _pptx_to_pdf(path, asset_dir / "render")
        if pdf:
            rendered = _render_pdf(pdf, asset_dir / "slides")
        for number, slide in enumerate(deck.slides, 1):
            text = "\n".join(shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip())
            locator = f"Slide {number}"
            image = str(rendered[number - 1]) if number <= len(rendered) else None
            chunks.append(SourceChunk(id=_id(path, locator, "text"), document=path.name, material=material,
                                      locator=locator, text=text or "[Visual slide with no extractable text]", image_path=image))
            if image:
                chunks.append(SourceChunk(id=_id(path, locator, "visual"), document=path.name, material=material,
                                          locator=locator, text=f"Visual evidence from {locator}: {text}", image_path=image, kind="visual"))
    elif path.suffix.lower() == ".pdf":
        chunks.extend(_ingest_pdf_pages(path, path.name, material, output_root))
    elif path.suffix.lower() in {".ppt", ".odp"}:
        pdf = _pptx_to_pdf(path, asset_dir / "render")
        if not pdf:
            raise RuntimeError(f"Could not convert {path.name}; LibreOffice is required")
        chunks.extend(_ingest_pdf_pages(pdf, path.name, material, output_root, locator_name="Slide"))
    elif path.suffix.lower() == ".docx":
        doc = Document(path)
        for number, para in enumerate((p for p in doc.paragraphs if p.text.strip()), 1):
            locator = f"Paragraph {number}"
            chunks.append(SourceChunk(id=_id(path, locator, "text"), document=path.name, material=material,
                                      locator=locator, text=para.text.strip()))
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        for number, paragraph in enumerate(paragraphs, 1):
            locator = f"Section {number}"
            chunks.append(SourceChunk(id=_id(path, locator, "text"), document=path.name, material=material,
                                      locator=locator, text=paragraph))
    return chunks


def build_index(inputs: list[Path], output_root: Path) -> list[SourceChunk]:
    output_root.mkdir(parents=True, exist_ok=True)
    chunks = [chunk for path in inputs for chunk in ingest_file(path, output_root)]
    target = output_root / "chunks.jsonl"
    with target.open("w", encoding="utf-8") as stream:
        for chunk in chunks:
            stream.write(chunk.model_dump_json() + "\n")
    return chunks


def load_index(path: Path) -> list[SourceChunk]:
    return [SourceChunk.model_validate_json(line) for line in path.read_text(encoding="utf-8").splitlines() if line]
